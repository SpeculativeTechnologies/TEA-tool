#!/usr/bin/env python3
"""Render a run's model outputs into a human-friendly Excel workbook.

`/tea-model` (dummy run) and `/tea-data` (real run) both call this *after*
`model.py` has produced its results JSON and plot PNGs. The workbook is a
**read-only presentation layer** for people who are more comfortable with
spreadsheets than with code or JSON:

- It contains **values only — never formulas**. Every number is computed by
  `model.py`; this script just lays it out. Editing a cell here changes nothing;
  to change results you edit the inputs and re-run the model.
- It mirrors the layout of a hand-built TEA workbook (see the reference
  `Techonomics - Polymerization Catalyst` model): an Overview/Documentation
  sheet, a Dashboard of headline numbers, a Main Input sheet with the input
  cells shaded grey, one sheet per block of intermediate model outputs, and a
  Charts sheet that embeds the rendered figures.

It is deliberately **generic**: it walks whatever `model-results.json` the run's
bespoke `model.py` emitted and renders each top-level section as a table, so it
works for any TEA model without per-run code.

Usage:
    python scripts/export_workbook.py --run-dir runs/<project> [--tag DUMMY]
    python scripts/export_workbook.py --inputs IN.csv --results R.json \
        --plots-dir DIR --output OUT.xlsx [--tag DUMMY] [--project "Name"]

Dependency: openpyxl (the plotting layer already needs matplotlib; the core
pipeline stays standard-library-only).
"""
from __future__ import annotations

import argparse
import csv
import json
import struct
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- palette (kept close to a conventional engineering workbook) ------------
NAVY = "1F3864"
MIDBLUE = "2E5395"
LIGHTBLUE = "D6E0F0"
GREY_INPUT = "D9D9D9"   # user-input cells, shaded grey by convention
BANNER_DUMMY = "C00000"
BANNER_REAL = "1F6E43"
WHITE = "FFFFFF"

TITLE_FONT = Font(bold=True, size=16, color=WHITE)
H_FONT = Font(bold=True, size=11, color=WHITE)
SECTION_FONT = Font(bold=True, size=12, color=NAVY)
LABEL_FONT = Font(bold=True, color=NAVY)
NOTE_FONT = Font(italic=True, size=9, color="595959")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")

# Result keys that are bookkeeping, not model outputs (handled specially / hidden).
META_KEYS = {"tag", "inputs_path", "plots"}

# Input-CSV column aliases (mirror model.py's load_inputs so the same files work).
ENTITY_KEYS = ("entity", "unit")
PARAM_KEYS = ("parameter", "field")
VALUE_KEYS = ("value", "normalized_value", "capex_usd")
UNITS_KEYS = ("normalized_units", "raw_units", "units")
SOURCE_KEYS = ("source_key", "source")
URL_KEYS = ("source_url", "url")
NOTE_KEYS = ("notes", "quality_tag")


# --- small helpers ----------------------------------------------------------
def _first(row: dict, keys) -> str | None:
    for key in keys:
        val = row.get(key)
        if val not in (None, ""):
            return val
    return None


def _num_format(value: float) -> str:
    a = abs(value)
    if value == int(value) and a < 1e15:
        return "#,##0"
    if a >= 1000:
        return "#,##0"
    if a >= 1:
        return "#,##0.00"
    if a > 0:
        return "0.0000"
    return "0"


def _set(ws, row, col, value, *, font=None, fill=None, border=False,
         align=None, number_format=None):
    cell = ws.cell(row=row, column=col, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if border:
        cell.border = BORDER
    if align:
        cell.alignment = align
    if number_format:
        cell.number_format = number_format
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.number_format = _num_format(float(value))
    return cell


def _autosize(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width, 8), 60)


def _track(widths: dict[int, int], col: int, value) -> None:
    length = len(str(value)) if value is not None else 0
    widths[col] = max(widths.get(col, 0), length + 2)


def _png_size(path: Path) -> tuple[int, int]:
    """Read a PNG's pixel dimensions from its IHDR (no Pillow dependency)."""
    try:
        with open(path, "rb") as fh:
            head = fh.read(24)
        if head[:8] == b"\x89PNG\r\n\x1a\n" and head[12:16] == b"IHDR":
            w, h = struct.unpack(">II", head[16:24])
            return int(w), int(h)
    except (OSError, struct.error):
        pass
    return 0, 0


def _pretty(text: str) -> str:
    return str(text).replace("_", " ").strip().capitalize()


# --- sheet builders ---------------------------------------------------------
def _banner(ws, title: str, subtitle: str, is_dummy: bool) -> int:
    ws.merge_cells("A1:F1")
    _set(ws, 1, 1, title, font=TITLE_FONT, fill=NAVY, align=Alignment(vertical="center", indent=1))
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:F2")
    flag = ("DUMMY DATA — structure check only, NOT real numbers" if is_dummy
            else "REAL DATA — citation-backed model run")
    _set(ws, 2, 1, flag, font=Font(bold=True, color=WHITE),
         fill=(BANNER_DUMMY if is_dummy else BANNER_REAL),
         align=Alignment(vertical="center", indent=1))
    ws.row_dimensions[2].height = 20
    if subtitle:
        ws.merge_cells("A3:F3")
        _set(ws, 3, 1, subtitle, font=NOTE_FONT, align=WRAP)
    return 5


def build_overview(wb, *, project, is_dummy, inputs_path, results, plots,
                   results_path):
    ws = wb.active
    ws.title = "Overview"
    ws.sheet_view.showGridLines = False
    row = _banner(ws, project, "Techno-economic model — generated workbook", is_dummy)

    _set(ws, row, 1, "About this workbook", font=SECTION_FONT)
    row += 1
    blurb = (
        "This workbook is generated automatically from the model code (model.py). "
        "It holds VALUES ONLY — there are no live spreadsheet formulas. Every number "
        "was computed in code and written here for easy review. To change any result, "
        "edit the model inputs and re-run the model; this file is then regenerated."
    )
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=6)
    _set(ws, row, 1, blurb, font=Font(size=10), align=WRAP)
    row += 4

    _set(ws, row, 1, "Run details", font=SECTION_FONT)
    row += 1
    headline = results.get("headline", {})
    details = [
        ("Data status", "DUMMY (placeholder)" if is_dummy else "Real (cited)"),
        ("Inputs file", inputs_path),
        ("Results file", results_path),
        ("Input cells", _count_inputs(inputs_path)),
        ("Headline metrics", len(_flatten_scalars(headline))),
        ("Charts", len(plots)),
        ("Output sections", sum(1 for k in results if k not in META_KEYS)),
    ]
    for label, value in details:
        _set(ws, row, 1, label, font=LABEL_FONT)
        _set(ws, row, 2, value)
        row += 1

    _autosize(ws, {1: 22, 2: 60})
    return ws


def build_documentation(wb, *, docs, is_dummy):
    """Embed the run's narrative (scope.md, model-requirements.md) as plain text,
    mirroring the Documentation sheet of a conventional TEA workbook."""
    docs = [(title, Path(p)) for title, p in (docs or []) if Path(p).exists()]
    if not docs:
        return None
    ws = wb.create_sheet("Documentation")
    ws.sheet_view.showGridLines = False
    row = _banner(ws, "Documentation", "Project scope and modelling assumptions.", is_dummy)
    ws.column_dimensions["A"].width = 110
    for title, path in docs:
        _set(ws, row, 1, title, font=SECTION_FONT)
        row += 1
        for line in path.read_text(encoding="utf-8").splitlines():
            stripped = line.strip()
            if not stripped or stripped.startswith(("<!--", "-->")):
                continue
            if stripped.startswith("#"):
                text = stripped.lstrip("# ").strip()
                _set(ws, row, 1, text, font=LABEL_FONT)
            else:
                _set(ws, row, 1, stripped, font=Font(size=10), align=WRAP)
            row += 1
        row += 2
    return ws


def build_dashboard(wb, *, results, is_dummy):
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    row = _banner(ws, "Dashboard — headline results", "", is_dummy)

    headline = results.get("headline", {})
    if not headline:
        _set(ws, row, 1, "(model emitted no headline block)", font=NOTE_FONT)
        return ws

    _set(ws, row, 1, "Metric", font=H_FONT, fill=MIDBLUE, border=True)
    _set(ws, row, 2, "Value", font=H_FONT, fill=MIDBLUE, border=True, align=CENTER)
    row += 1
    widths: dict[int, int] = {1: 30, 2: 16}
    for path, value in _flatten_scalars(headline):
        label = " — ".join(_pretty(p) for p in path)
        _set(ws, row, 1, label, border=True)
        _track(widths, 1, label)
        if isinstance(value, bool):
            _set(ws, row, 2, "PASS" if value else "FAIL", border=True, align=CENTER,
                 font=Font(bold=True, color=(BANNER_REAL if value else BANNER_DUMMY)))
        else:
            _set(ws, row, 2, value, border=True, align=Alignment(horizontal="right"))
        _track(widths, 2, value)
        row += 1
    _autosize(ws, widths)
    return ws


def build_main_input(wb, *, inputs_path, is_dummy):
    ws = wb.create_sheet("Main Input")
    ws.sheet_view.showGridLines = False
    row = _banner(
        ws, "Main Input", "The values fed into the model. Input cells are shaded grey.",
        is_dummy,
    )

    rows = _read_input_rows(inputs_path)
    cols = [
        ("Entity", ENTITY_KEYS, False),
        ("Parameter", PARAM_KEYS, False),
        ("Value", VALUE_KEYS, True),       # shaded grey = the editable input
        ("Units", UNITS_KEYS, False),
        ("Source", SOURCE_KEYS, False),
        ("Source URL", URL_KEYS, False),
        ("Notes", NOTE_KEYS, False),
    ]
    # Drop entirely-empty optional columns (keeps the dummy CSV view tidy).
    present = [c for c in cols if c[0] in ("Entity", "Parameter", "Value")
               or any(_first(r, c[1]) for r in rows)]

    widths: dict[int, int] = {}
    for ci, (label, _keys, _is_input) in enumerate(present, start=1):
        _set(ws, row, ci, label, font=H_FONT, fill=MIDBLUE, border=True, align=CENTER)
        _track(widths, ci, label)
    header_row = row
    row += 1
    for r in rows:
        for ci, (_label, keys, is_input) in enumerate(present, start=1):
            raw = _first(r, keys)
            value = _coerce(raw)
            fill = GREY_INPUT if is_input else None
            _set(ws, row, ci, value, border=True, fill=fill, align=WRAP)
            _track(widths, ci, raw)
        row += 1
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    _autosize(ws, widths)
    return ws


def build_output_sheets(wb, *, results, is_dummy):
    """One sheet per top-level results section (the intermediate model outputs)."""
    made = []
    used_titles: set[str] = set()
    for key, value in results.items():
        if key in META_KEYS or key == "headline":
            continue
        title = _sheet_title(_pretty(key), used_titles)
        ws = wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        row = _banner(ws, _pretty(key), "Intermediate model output.", is_dummy)
        widths: dict[int, int] = {}
        _render_block(ws, row, value, widths)
        _autosize(ws, widths)
        made.append(ws)
    return made


def build_charts(wb, *, plots, is_dummy):
    if not plots:
        return None
    ws = wb.create_sheet("Charts")
    ws.sheet_view.showGridLines = False
    _banner(ws, "Charts", "Figures rendered by the model.", is_dummy)
    anchor_row = 5
    ws.column_dimensions["A"].width = 12
    for path in plots:
        p = Path(path)
        if not p.exists():
            continue
        _set(ws, anchor_row, 1, p.stem.replace("_", " "), font=SECTION_FONT)
        anchor_row += 1
        try:
            img = XLImage(str(p))
            w, h = _png_size(p)
            if w and h:
                target_w = 720
                img.width = target_w
                img.height = int(h * target_w / w)
            ws.add_image(img, f"A{anchor_row}")
            rows_spanned = int((img.height or 360) / 18) + 2
        except Exception:  # noqa: BLE001 — never let one bad PNG kill the workbook
            _set(ws, anchor_row, 1, f"(could not embed {p.name})", font=NOTE_FONT)
            rows_spanned = 2
        anchor_row += rows_spanned + 2
    return ws


# --- results flattening / rendering -----------------------------------------
def _flatten_scalars(obj, prefix=()):
    """Yield (path_tuple, scalar) for every scalar leaf in a nested dict."""
    out = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            out.extend(_flatten_scalars(v, prefix + (k,)))
    elif isinstance(obj, (list, tuple)):
        # Render a short scalar list inline; skip nested lists.
        if obj and all(not isinstance(x, (dict, list)) for x in obj):
            out.append((prefix, ", ".join(str(x) for x in obj)))
    elif obj is not None:
        out.append((prefix, obj))
    return out


def _is_scalar(v):
    return v is None or isinstance(v, (int, float, str, bool))


def _render_block(ws, row, value, widths) -> int:
    """Render an arbitrary results section as readable table(s)."""
    if _is_scalar(value):
        _set(ws, row, 1, value, border=True)
        _track(widths, 1, value)
        return row + 2

    if isinstance(value, list):
        return _render_list(ws, row, value, widths)

    if isinstance(value, dict):
        # dict[str, scalar] -> two-column table
        if all(_is_scalar(v) or _is_short_list(v) for v in value.values()):
            return _render_keyvalue(ws, row, value, widths)
        # dict[str, dict[scalar]] -> matrix (columns are the outer keys)
        if all(isinstance(v, dict) and all(_is_scalar(x) for x in v.values())
               for v in value.values()):
            return _render_matrix(ws, row, value, widths)
        # mixed / deeper -> one labelled sub-table per outer key
        for sub_key, sub_val in value.items():
            _set(ws, row, 1, _pretty(sub_key), font=SECTION_FONT)
            row += 1
            row = _render_block(ws, row, sub_val, widths)
            row += 1
        return row
    # fallback
    _set(ws, row, 1, json.dumps(value), align=WRAP)
    return row + 2


def _is_short_list(v):
    return isinstance(v, list) and all(not isinstance(x, (dict, list)) for x in v)


def _render_keyvalue(ws, row, mapping, widths) -> int:
    _set(ws, row, 1, "Item", font=H_FONT, fill=MIDBLUE, border=True)
    _set(ws, row, 2, "Value", font=H_FONT, fill=MIDBLUE, border=True, align=CENTER)
    row += 1
    for k, v in mapping.items():
        label = _pretty(k)
        _set(ws, row, 1, label, border=True)
        _track(widths, 1, label)
        if _is_short_list(v):
            v = ", ".join(str(x) for x in v)
        if isinstance(v, bool):
            _set(ws, row, 2, "PASS" if v else "FAIL", border=True, align=CENTER)
        else:
            _set(ws, row, 2, v, border=True, align=Alignment(horizontal="right"))
        _track(widths, 2, v)
        row += 1
    return row + 1


def _render_matrix(ws, row, mapping, widths) -> int:
    outer = list(mapping.keys())
    inner: list[str] = []
    for v in mapping.values():
        for k in v:
            if k not in inner:
                inner.append(k)
    _set(ws, row, 1, "", font=H_FONT, fill=MIDBLUE, border=True)
    for ci, o in enumerate(outer, start=2):
        label = _pretty(o)
        _set(ws, row, ci, label, font=H_FONT, fill=MIDBLUE, border=True, align=CENTER)
        _track(widths, ci, label)
    row += 1
    for ik in inner:
        label = _pretty(ik)
        _set(ws, row, 1, label, font=LABEL_FONT, fill=LIGHTBLUE, border=True)
        _track(widths, 1, label)
        for ci, o in enumerate(outer, start=2):
            v = mapping[o].get(ik)
            if isinstance(v, bool):
                _set(ws, row, ci, "PASS" if v else "FAIL", border=True, align=CENTER)
            else:
                _set(ws, row, ci, v, border=True, align=Alignment(horizontal="right"))
            _track(widths, ci, v)
        row += 1
    return row + 1


def _render_list(ws, row, items, widths) -> int:
    if items and all(isinstance(x, dict) for x in items):
        cols: list[str] = []
        for it in items:
            for k in it:
                if k not in cols:
                    cols.append(k)
        for ci, c in enumerate(cols, start=1):
            label = _pretty(c)
            _set(ws, row, ci, label, font=H_FONT, fill=MIDBLUE, border=True, align=CENTER)
            _track(widths, ci, label)
        row += 1
        for it in items:
            for ci, c in enumerate(cols, start=1):
                _set(ws, row, ci, it.get(c), border=True)
                _track(widths, ci, it.get(c))
            row += 1
        return row + 1
    for x in items:
        _set(ws, row, 1, x, border=True)
        _track(widths, 1, x)
        row += 1
    return row + 1


# --- input reading ----------------------------------------------------------
def _read_input_rows(path) -> list[dict]:
    p = Path(path)
    if not p.exists():
        return []
    with open(p, newline="", encoding="utf-8") as fh:
        lines = [ln for ln in fh if not ln.lstrip().startswith("#")]
    return list(csv.DictReader(lines))


def _count_inputs(path) -> int:
    rows = _read_input_rows(path)
    return sum(1 for r in rows if _first(r, VALUE_KEYS) not in (None, ""))


def _coerce(raw):
    if raw in (None, ""):
        return None
    try:
        return float(str(raw).replace(",", ""))
    except ValueError:
        return raw


def _sheet_title(name: str, used: set[str]) -> str:
    base = name[:31] or "Sheet"
    title = base
    n = 2
    while title in used:
        suffix = f" ({n})"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


# --- entrypoint -------------------------------------------------------------
def build_workbook(*, inputs_path, results_path, plots_dir, output_path,
                   tag="", project=None, docs=None):
    results = json.loads(Path(results_path).read_text(encoding="utf-8"))
    is_dummy = (tag or results.get("tag") or "").upper() == "DUMMY"
    project = project or Path(output_path).stem.replace("-", " ").replace("_", " ")

    # Prefer the plot list recorded in the results JSON; fall back to the dir.
    plots = [Path(p) for p in results.get("plots", []) if Path(p).exists()]
    if not plots and plots_dir and Path(plots_dir).is_dir():
        plots = sorted(Path(plots_dir).glob("*.png"))

    wb = Workbook()
    build_overview(wb, project=project, is_dummy=is_dummy, inputs_path=str(inputs_path),
                   results=results, plots=plots, results_path=str(results_path))
    build_documentation(wb, docs=docs, is_dummy=is_dummy)
    build_dashboard(wb, results=results, is_dummy=is_dummy)
    build_main_input(wb, inputs_path=inputs_path, is_dummy=is_dummy)
    build_output_sheets(wb, results=results, is_dummy=is_dummy)
    build_charts(wb, plots=plots, is_dummy=is_dummy)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {"output": str(output_path).replace("\\", "/"), "dummy": is_dummy,
            "sheets": wb.sheetnames, "charts": len(plots)}


def _defaults_from_run(run_dir: Path, tag: str):
    is_dummy = (tag or "").upper() == "DUMMY"
    inputs = run_dir / ("model/dummy-inputs.csv" if is_dummy else "Main Input.csv")
    results = run_dir / ("model-results-DUMMY.json" if is_dummy else "model-results.json")
    plots = run_dir / ("plots-dummy" if is_dummy else "plots")
    out = run_dir / (f"{run_dir.name}-model-workbook-DUMMY.xlsx" if is_dummy
                     else f"{run_dir.name}-model-workbook.xlsx")
    return inputs, results, plots, out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--run-dir", help="run directory; fills in the defaults below")
    ap.add_argument("--inputs", help="inputs CSV (dummy-inputs.csv or Main Input.csv)")
    ap.add_argument("--results", help="model-results.json")
    ap.add_argument("--plots-dir", help="directory of rendered PNGs")
    ap.add_argument("--output", help="output .xlsx path")
    ap.add_argument("--tag", default="", help="DUMMY for a dry-run workbook; blank for real")
    ap.add_argument("--project", help="display name (default: derived)")
    args = ap.parse_args(argv)

    docs = None
    if args.run_dir:
        run_dir = Path(args.run_dir)
        d_in, d_res, d_plots, d_out = _defaults_from_run(run_dir, args.tag)
        inputs = Path(args.inputs) if args.inputs else d_in
        results = Path(args.results) if args.results else d_res
        plots = Path(args.plots_dir) if args.plots_dir else d_plots
        output = Path(args.output) if args.output else d_out
        docs = [("Scope", run_dir / "scope.md"),
                ("Model requirements", run_dir / "model-requirements.md")]
    else:
        missing = [n for n in ("inputs", "results", "output") if not getattr(args, n)]
        if missing:
            ap.error(f"--{', --'.join(missing)} required when --run-dir is not given")
        inputs, results, output = Path(args.inputs), Path(args.results), Path(args.output)
        plots = Path(args.plots_dir) if args.plots_dir else None

    summary = build_workbook(
        inputs_path=inputs, results_path=results, plots_dir=plots,
        output_path=output, tag=args.tag, project=args.project, docs=docs,
    )
    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
