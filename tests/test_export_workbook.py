"""Tests for the Excel workbook exporter (scripts/export_workbook.py).

The workbook is a read-only, values-only presentation layer over a run's
model-results.json + inputs CSV + plot PNGs. These tests prove it:

- builds the expected sheets (Overview, Dashboard, Main Input, one per output
  section, Charts),
- carries values only — no formulas leak into any cell,
- labels a dummy run as DUMMY, and
- shades the Main Input value cells grey (the spreadsheet-convention input cell).
"""

import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

try:
    import openpyxl  # noqa: F401
    import export_workbook
    HAVE_OPENPYXL = True
except ImportError:  # pragma: no cover - exercised only without the optional dep
    HAVE_OPENPYXL = False

RESULTS = {
    "headline": {
        "total_cost_per_fu": 1.65,
        "targets_scorecard": {"cost_reduction": True, "energy": False},
    },
    "breakdown": {"compute": 1.2, "power": 0.45},
    "groups": {
        "route_a": {"energy": 10.0, "cost": 2.0},
        "route_b": {"energy": 7.5, "cost": 1.5},
    },
    "tag": "DUMMY",
}

INPUTS = (
    "entity,parameter,value\n"
    "compute,cost_per_fu,1.20\n"
    "power,cost_per_fu,0.45\n"
)


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class ExportWorkbookTests(unittest.TestCase):
    def _build(self, tag="DUMMY"):
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: __import__("shutil").rmtree(tmp, ignore_errors=True))
        inputs = tmp / "inputs.csv"
        inputs.write_text(INPUTS, encoding="utf-8")
        results = dict(RESULTS, tag=tag)
        results_path = tmp / "model-results.json"
        results_path.write_text(json.dumps(results), encoding="utf-8")
        out = tmp / "wb.xlsx"
        summary = export_workbook.build_workbook(
            inputs_path=inputs, results_path=results_path, plots_dir=None,
            output_path=out, tag=tag, project="Test Project",
        )
        return out, summary

    def test_expected_sheets(self):
        out, summary = self._build()
        wb = openpyxl.load_workbook(out)
        for name in ("Overview", "Dashboard", "Main Input", "Breakdown", "Groups"):
            self.assertIn(name, wb.sheetnames, summary)

    def test_no_formulas_anywhere(self):
        out, _ = self._build()
        wb = openpyxl.load_workbook(out)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        self.assertFalse(
                            cell.value.startswith("="),
                            f"formula leaked into {ws.title}!{cell.coordinate}",
                        )

    def test_dummy_banner(self):
        out, summary = self._build(tag="DUMMY")
        self.assertTrue(summary["dummy"])
        wb = openpyxl.load_workbook(out)
        banner = wb["Dashboard"]["A2"].value
        self.assertIn("DUMMY", banner)

    def test_real_banner(self):
        out, summary = self._build(tag="")
        self.assertFalse(summary["dummy"])
        wb = openpyxl.load_workbook(out)
        self.assertIn("REAL", wb["Dashboard"]["A2"].value)

    def test_dashboard_carries_headline_values(self):
        out, _ = self._build()
        wb = openpyxl.load_workbook(out)
        vals = [c.value for row in wb["Dashboard"].iter_rows() for c in row]
        self.assertIn(1.65, vals)

    def test_main_input_value_cells_shaded(self):
        out, _ = self._build()
        wb = openpyxl.load_workbook(out)
        ws = wb["Main Input"]
        # Find the header row, then confirm the Value column data cells are grey.
        header_row = next(
            r[0].row for r in ws.iter_rows()
            if any(c.value == "Value" for c in r)
        )
        value_col = next(
            c.column for c in ws[header_row] if c.value == "Value"
        )
        data_cell = ws.cell(row=header_row + 1, column=value_col)
        self.assertEqual(data_cell.value, 1.20)
        self.assertEqual((data_cell.fill.fgColor.rgb or "")[-6:], export_workbook.GREY_INPUT)


if __name__ == "__main__":
    unittest.main()
